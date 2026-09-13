"""Fast default tests use SQLite; TEST_DATABASE_URL enables real PostgreSQL.

Only a database whose name ends in _test is accepted. Each run creates and
removes its own schema; application tables outside that schema are untouched.
"""
import os
import secrets
from uuid import uuid4

# Set isolated configuration before importing application modules; never use .env secrets.
os.environ["APP_ENV"] = "test"
os.environ["APP_NAME"] = "Legal AI Auth Tests"
os.environ["SECRET_KEY"] = secrets.token_urlsafe(48)
os.environ["ALGORITHM"] = "HS256"
os.environ["ACCESS_TOKEN_EXPIRE_MINUTES"] = "30"
os.environ["CHUNK_SIZE"] = "1200"
os.environ["CHUNK_OVERLAP"] = "200"
os.environ["DATABASE_URL"] = os.environ.get(
    "TEST_DATABASE_URL", "postgresql+psycopg://localhost:1/legal_ai_test"
)

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import main
from app.core.config import PROJECT_ROOT, get_settings
from app.db.session import get_db


@pytest.fixture(scope="session")
def db_engine(tmp_path_factory):
    database_url = os.environ.get("TEST_DATABASE_URL")
    admin_engine = None
    schema = None
    if database_url:
        url = make_url(database_url)
        if url.drivername != "postgresql+psycopg" or not (url.database or "").endswith("_test"):
            raise ValueError("TEST_DATABASE_URL must use psycopg and a database ending in _test")
        schema = f"auth_test_{uuid4().hex}"
        admin_engine = create_engine(url)
        with admin_engine.begin() as connection:
            connection.execute(CreateSchema(schema))
        engine = create_engine(url, connect_args={"options": f"-csearch_path={schema}"})
    else:
        path = tmp_path_factory.mktemp("auth") / "auth.db"
        engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
        @event.listens_for(engine, "connect")
        def enable_foreign_keys(connection, record):
            connection.execute("PRAGMA foreign_keys=ON")

    try:
        # Exercise the actual migration, never Base.metadata.create_all().
        config = Config(str(PROJECT_ROOT / "alembic.ini"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "head")
        yield engine
    finally:
        engine.dispose()
        if admin_engine is not None:
            try:
                with admin_engine.begin() as connection:
                    connection.execute(DropSchema(schema, cascade=True))
            finally:
                admin_engine.dispose()


@pytest.fixture
def db_session(db_engine):
    with db_engine.connect() as connection:
        # Explicit SQL BEGIN avoids SQLite's legacy SAVEPOINT transaction behavior.
        if db_engine.dialect.name == "sqlite":
            connection.exec_driver_sql("BEGIN")
            transaction = connection.get_transaction()
        else:
            transaction = connection.begin()
        with Session(bind=connection, join_transaction_mode="create_savepoint") as session:
            try:
                yield session
            finally:
                session.close()
                transaction.rollback()


@pytest.fixture
def document_storage(tmp_path, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "document_storage_path", tmp_path / "documents")
    monkeypatch.setattr(settings, "max_upload_size_mb", 1)
    return settings.storage_directory


@pytest.fixture
def client(db_session, db_engine, monkeypatch, document_storage):
    def override_get_db():
        yield db_session

    main.app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(main, "engine", db_engine)
    try:
        with TestClient(main.app) as test_client:
            yield test_client
    finally:
        main.app.dependency_overrides.clear()


@pytest.fixture
def docx_bytes():
    from io import BytesIO
    from docx import Document

    document = Document()
    document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    document.add_paragraph("Khoản 2. Quyền và nghĩa vụ.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def pdf_bytes():
    from io import BytesIO
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    for sentence in [b"Article 1. Scope.", b"Article 2. Duties."]:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 40 240 Td (" + sentence + b") Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture
def blank_pdf_bytes():
    from io import BytesIO
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
